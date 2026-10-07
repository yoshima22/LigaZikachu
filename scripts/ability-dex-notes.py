"""Acrescenta à planilha a coluna 'Na Pokédex (resumo)' e renomeia efeitos genéricos.

Os efeitos compartilhados por várias habilidades (ex.: Corpo limpo) tinham o nome de UMA delas.
Aqui cada habilidade ganha uma linha dizendo o que ela faz na Pokédex, e o efeito ganha um nome neutro.
Uso: python scripts/ability-dex-notes.py   (depois: python scripts/generate-ability-data.py)
"""
import openpyxl

SRC = 'docs/habilidades-passivas-mascotes.xlsx'

RENAME = {
    'CORPO_LIMPO': 'Resistência a debuff',
    'CASCO': 'Defesa leve',
    'PODER_BRUTO': 'Ataque reforçado',
    'QUEBRA_GUARDA': 'Golpe certeiro',
}

# slug -> o que a habilidade faz na Pokédex (resumo em português)
NOTES = {
    # Resistência a debuff
    'clear-body': 'Impede que outros Pokémon reduzam seus atributos.',
    'white-smoke': 'Impede que outros Pokémon reduzam seus atributos.',
    'full-metal-body': 'Impede que outros Pokémon reduzam seus atributos.',
    'hyper-cutter': 'Impede que o Ataque seja reduzido.',
    'big-pecks': 'Impede que a Defesa seja reduzida.',
    'keen-eye': 'Impede que a precisão seja reduzida.',
    'shield-dust': 'Bloqueia os efeitos extras de golpes recebidos.',
    'inner-focus': 'Impede recuar (flinch) e, nas gerações novas, também a Intimidação.',
    'immunity': 'Impede envenenamento.',
    'insomnia': 'Impede o sono.',
    'vital-spirit': 'Impede o sono.',
    'limber': 'Impede paralisia.',
    'magma-armor': 'Impede congelamento.',
    'own-tempo': 'Impede confusão.',
    'oblivious': 'Impede paixão e provocação.',
    'water-veil': 'Impede queimaduras.',
    'leaf-guard': 'Protege de problemas de status sob sol forte.',
    'soundproof': 'Bloqueia golpes sonoros.',
    'purifying-salt': 'Protege de problemas de status e reduz dano Fantasma.',
    'good-as-gold': 'Imune a golpes de status.',
    # Defesa leve
    'shell-armor': 'Impede golpes críticos.',
    'battle-armor': 'Impede golpes críticos.',
    'rock-head': 'Impede o dano de recuo dos próprios golpes.',
    'overcoat': 'Protege de dano de clima e de golpes de pó.',
    'bulletproof': 'Bloqueia golpes de bomba, bola e projétil.',
    'fluffy': 'Reduz pela metade o dano de contato (mas dobra o de Fogo).',
    'fur-coat': 'Reduz pela metade o dano físico.',
    'ice-scales': 'Reduz pela metade o dano especial.',
    'aura-guard': 'Reduz pela metade o dano de golpes de contato.',
    # Golpe certeiro
    'compound-eyes': 'Aumenta a precisão dos golpes.',
    'victory-star': 'Aumenta a precisão da equipe.',
    'minds-eye': 'Ignora a esquiva do alvo e não perde precisão.',
    'no-guard': 'Todos os golpes (dele e contra ele) acertam.',
    'mold-breaker': 'Ignora habilidades do alvo que atrapalhariam o golpe.',
    'teravolt': 'Ignora habilidades do alvo que atrapalhariam o golpe.',
    'turboblaze': 'Ignora habilidades do alvo que atrapalhariam o golpe.',
    'infiltrator': 'Atravessa barreiras como Reflect e Light Screen.',
    'scrappy': 'Golpes Normais e Lutadores acertam Fantasmas.',
    'sniper': 'Golpes críticos causam ainda mais dano.',
    'super-luck': 'Aumenta a chance de golpe crítico.',
    'merciless': 'Golpes sempre críticos contra alvos envenenados.',
    'stakeout': 'Dobra o dano contra quem acabou de entrar.',
    'unseen-fist': 'Golpes de contato atravessam a proteção do alvo.',
    'piercing-drill': 'Golpes de contato atravessam a proteção do alvo.',
    'long-reach': 'Seus golpes nunca fazem contato.',
    # Ataque reforçado
    'huge-power': 'Dobra o Ataque.',
    'pure-power': 'Dobra o Ataque.',
    'hustle': 'Golpes físicos mais fortes, porém menos precisos.',
    'gorilla-tactics': 'Ataque maior, mas fica preso ao primeiro golpe escolhido.',
    'intrepid-sword': 'Aumenta o Ataque ao entrar em combate.',
    'solar-power': 'Ataque Especial maior sob sol forte, com custo de HP.',
    'protosynthesis': 'Aumenta o maior atributo sob sol forte.',
    'quark-drive': 'Aumenta o maior atributo em terreno elétrico.',
    # Maestria
    'technician': 'Golpes fracos (até 60 de poder) ficam 50% mais fortes.',
    'iron-fist': 'Golpes de soco ficam mais fortes.',
    'strong-jaw': 'Golpes de mordida ficam mais fortes.',
    'tough-claws': 'Golpes de contato ficam mais fortes.',
    'sharpness': 'Golpes cortantes ficam mais fortes.',
    'mega-launcher': 'Golpes de pulso e aura ficam mais fortes.',
    'punk-rock': 'Golpes sonoros ficam mais fortes e causam menos dano a ele.',
    'reckless': 'Golpes com recuo ficam mais fortes.',
    'sheer-force': 'Golpes com efeito extra ficam mais fortes, mas perdem o efeito.',
    'steelworker': 'Golpes de Aço ficam mais fortes.',
    'transistor': 'Golpes Elétricos ficam mais fortes.',
    'dragons-maw': 'Golpes de Dragão ficam mais fortes.',
    'rocky-payload': 'Golpes de Pedra ficam mais fortes.',
    'fire-mane': 'Golpes de Fogo ficam mais fortes.',
    'sand-force': 'Golpes de Pedra, Terra e Aço ficam mais fortes na tempestade de areia.',
    'liquid-voice': 'Golpes sonoros viram golpes de Água.',
    'aerilate': 'Golpes Normais viram Voadores, mais fortes.',
    'pixilate': 'Golpes Normais viram Fada, mais fortes.',
    'refrigerate': 'Golpes Normais viram Gelo, mais fortes.',
    'dragonize': 'Golpes Normais viram Dragão, mais fortes.',
    # Impulso
    'chlorophyll': 'Dobra a Velocidade sob sol forte.',
    'swift-swim': 'Dobra a Velocidade na chuva.',
    'sand-rush': 'Dobra a Velocidade na tempestade de areia.',
    'slush-rush': 'Dobra a Velocidade no granizo.',
    'speed-boost': 'A Velocidade sobe um pouco a cada turno.',
    'unburden': 'Dobra a Velocidade ao usar ou perder o item.',
    'prankster': 'Golpes de status agem primeiro.',
    'gale-wings': 'Golpes Voadores agem primeiro.',
    'triage': 'Golpes de cura agem primeiro.',
    'moody': 'Todo turno sobe um atributo e reduz outro.',
    'rattled': 'Ganha Velocidade ao ser atingido por golpes Sombrios, Fantasmas ou de Inseto.',
    'weak-armor': 'Ao ser atingido por golpe físico, perde Defesa e ganha Velocidade.',
    # Cura
    'rain-dish': 'Recupera um pouco de HP todo turno na chuva.',
    'ice-body': 'Recupera um pouco de HP todo turno no granizo.',
    'dry-skin': 'Recupera HP na chuva e perde sob sol forte.',
    'poison-heal': 'Recupera HP todo turno quando envenenado.',
    'cheek-pouch': 'Recupera HP ao comer uma fruta.',
    'cud-chew': 'Reaproveita uma fruta já comida.',
    'harvest': 'Pode recuperar a fruta usada.',
    # Purificação
    'natural-cure': 'Cura problemas de status ao sair de campo.',
    'shed-skin': 'Chance de curar status todo turno.',
    'hydration': 'Cura status todo turno na chuva.',
    'healer': 'Chance de curar o status de aliados todo turno.',
    # Reforço / aura
    'battery': 'Aumenta o poder dos golpes especiais dos aliados.',
    'power-spot': 'Aumenta o poder dos golpes de quem está ao lado.',
    'plus': 'Aumenta o Ataque Especial junto de um aliado Plus/Minus.',
    'minus': 'Aumenta o Ataque Especial junto de um aliado Plus/Minus.',
    'flower-gift': 'Aumenta Ataque e Defesa Especial dos aliados sob sol forte.',
    'costar': 'Copia as mudanças de atributo de um aliado ao entrar.',
    'symbiosis': 'Passa o item para um aliado que ficou sem.',
    'steely-spirit': 'Golpes de Aço dos aliados ficam mais fortes.',
    'dark-aura': 'Golpes Sombrios de todos ficam mais fortes.',
    'fairy-aura': 'Golpes de Fada de todos ficam mais fortes.',
    # Véu
    'aroma-veil': 'Protege aliados de efeitos que afetam a mente.',
    'flower-veil': 'Protege aliados de Planta de perder atributos.',
    'sweet-veil': 'Impede os aliados de dormir.',
    # Pelagem / absorção
    'thick-fat': 'Reduz pela metade o dano de Fogo e Gelo.',
    'heatproof': 'Reduz pela metade o dano de Fogo e queimaduras.',
    'water-bubble': 'Reduz o dano de Fogo, dobra o de Água e impede queimaduras.',
    'levitate': 'Flutua: imune a golpes de Terra.',
    'eelevate': 'Flutua: imune a golpes de Terra.',
    'earth-eater': 'Recupera HP ao ser atingido por golpes de Terra.',
    'water-absorb': 'Recupera HP ao ser atingido por golpes de Água.',
    'volt-absorb': 'Recupera HP ao ser atingido por golpes Elétricos.',
    'flash-fire': 'Imune a Fogo e, ao absorver, fortalece os próprios golpes de Fogo.',
    'sap-sipper': 'Absorve golpes de Planta e ganha Ataque.',
    'motor-drive': 'Absorve golpes Elétricos e ganha Velocidade.',
    'well-baked-body': 'Imune a Fogo e ganha muita Defesa.',
    'wind-rider': 'Imune a golpes de vento e ganha Ataque.',
    # Reações e Controle
    'static': '30% de chance de paralisar quem faz contato.',
    'flame-body': '30% de chance de queimar quem faz contato.',
    'poison-point': '30% de chance de envenenar quem faz contato.',
    'effect-spore': '30% de chance de causar paralisia, veneno ou sono em quem faz contato.',
    'cute-charm': '30% de chance de causar paixão em quem faz contato.',
    'cursed-body': '30% de chance de bloquear o golpe que o atingiu.',
    'gooey': 'Reduz a Velocidade de quem faz contato.',
    'cotton-down': 'Reduz a Velocidade de todos quando atingido.',
    'mummy': 'Troca a habilidade de quem faz contato por Mummy.',
    'wandering-spirit': 'Troca de habilidade com quem faz contato.',
    'lingering-aroma': 'Troca a habilidade de quem faz contato.',
    'poison-touch': '30% de chance de envenenar ao atacar com contato.',
    'toxic-chain': 'Pode envenenar gravemente ao atingir um rival.',
    'stench': '10% de chance de fazer o alvo recuar.',
    'corrosion': 'Pode envenenar Pokémon de Aço e Veneno.',
    'poison-puppeteer': 'Alvos envenenados também ficam confusos.',
    'intimidate': 'Reduz o Ataque dos rivais ao entrar em combate.',
    'pressure': 'Faz os golpes contra ele gastarem mais PP.',
    'unnerve': 'Impede os rivais de comer frutas.',
    'dazzling': 'Impede os rivais de usar golpes com prioridade.',
    'queenly-majesty': 'Impede os rivais de usar golpes com prioridade.',
    'armor-tail': 'Impede os rivais de usar golpes com prioridade.',
    'supersweet-syrup': 'Reduz a esquiva dos rivais ao entrar.',
    'sword-of-ruin': 'Reduz a Defesa de todos os outros Pokémon.',
    'tablets-of-ruin': 'Reduz o Ataque de todos os outros Pokémon.',
    'vessel-of-ruin': 'Reduz o Ataque Especial de todos os outros Pokémon.',
    'beads-of-ruin': 'Reduz a Defesa Especial de todos os outros Pokémon.',
    # Dano condicional
    'guts': 'Ataque maior quando está com problema de status.',
    'toxic-boost': 'Ataque maior quando envenenado.',
    'flare-boost': 'Ataque Especial maior quando queimado.',
    'quick-feet': 'Velocidade maior quando está com problema de status.',
    'defiant': 'Ataque sobe muito ao ter um atributo reduzido.',
    'competitive': 'Ataque Especial sobe muito ao ter um atributo reduzido.',
    'guard-dog': 'Ataque sobe se for intimidado.',
    'justified': 'Ataque sobe ao ser atingido por golpe Sombrio.',
    'steam-engine': 'Velocidade sobe muito ao ser atingido por Fogo ou Água.',
    'thermal-exchange': 'Ataque sobe ao ser atingido por Fogo e não queima.',
    'electromorphosis': 'Ao ser atingido, o próximo golpe Elétrico dobra de poder.',
    'wind-power': 'Ao ser atingido por vento, o próximo golpe Elétrico dobra de poder.',
    'moxie': 'Ataque sobe ao derrotar um Pokémon.',
    'beast-boost': 'O maior atributo sobe ao derrotar um Pokémon.',
    'chilling-neigh': 'Ataque sobe ao derrotar um Pokémon.',
    'grim-neigh': 'Ataque Especial sobe ao derrotar um Pokémon.',
    'battle-bond': 'Transforma-se ao derrotar um Pokémon.',
    'soul-heart': 'Ataque Especial sobe sempre que qualquer Pokémon cai.',
    'receiver': 'Ao cair um aliado, ganha a habilidade dele.',
    'supreme-overlord': 'Ataque e Ataque Especial sobem por aliado derrotado.',
    'analytic': 'Golpes mais fortes se agir por último.',
    'download': 'Aumenta o ataque que explora a defesa mais fraca do rival.',
    'adaptability': 'O bônus de mesmo tipo (STAB) passa de 1,5× para 2×.',
    'neuroforce': 'Golpes super efetivos causam mais dano.',
    'tinted-lens': 'Golpes pouco efetivos causam o dobro de dano.',
    'parental-bond': 'Golpes batem duas vezes (a segunda com metade da força).',
    'skill-link': 'Golpes de vários acertos sempre acertam o máximo.',
    'anger-point': 'Ataque vai ao máximo ao receber um golpe crítico.',
    'berserk': 'Ataque Especial sobe quando o HP cai abaixo da metade.',
    'anger-shell': 'Com HP abaixo da metade, ganha ataque e perde defesa.',
    'blaze': 'Golpes de Fogo ficam mais fortes com HP baixo.',
    'torrent': 'Golpes de Água ficam mais fortes com HP baixo.',
    'overgrow': 'Golpes de Planta ficam mais fortes com HP baixo.',
    'swarm': 'Golpes de Inseto ficam mais fortes com HP baixo.',
    # Defesa
    'multiscale': 'Com HP cheio, reduz pela metade o dano recebido.',
    'shadow-shield': 'Com HP cheio, reduz pela metade o dano recebido.',
    'disguise': 'Um disfarce absorve o primeiro golpe.',
    'ice-face': 'Um bloco de gelo absorve o primeiro golpe físico.',
    'filter': 'Reduz o dano de golpes super efetivos.',
    'solid-rock': 'Reduz o dano de golpes super efetivos.',
    'prism-armor': 'Reduz o dano de golpes super efetivos.',
    'tera-shell': 'Com HP cheio, nenhum golpe é super efetivo contra ele.',
    'wonder-guard': 'Só golpes super efetivos o machucam.',
    'unaware': 'Ignora as mudanças de atributo do rival.',
    'sand-veil': 'Aumenta a esquiva na tempestade de areia.',
    'snow-cloak': 'Aumenta a esquiva no granizo.',
    'tangled-feet': 'Dobra a esquiva quando confuso.',
    'wonder-skin': 'Golpes de status contra ele erram mais.',
    'friend-guard': 'Reduz o dano que os aliados recebem.',
    'stamina': 'Defesa sobe ao ser atingido.',
    'water-compaction': 'Defesa sobe muito ao ser atingido por Água.',
    'dauntless-shield': 'Aumenta a Defesa ao entrar em combate.',
    'grass-pelt': 'Aumenta a Defesa em terreno de grama.',
    'marvel-scale': 'Defesa maior quando está com problema de status.',
    'emergency-exit': 'Sai de campo quando o HP cai abaixo da metade.',
    'wimp-out': 'Sai de campo quando o HP cai abaixo da metade.',
    'lightning-rod': 'Atrai golpes Elétricos e ganha Ataque Especial.',
    'storm-drain': 'Atrai golpes de Água e ganha Ataque Especial.',
    'magic-guard': 'Só sofre dano direto de golpes.',
    'magic-bounce': 'Reflete golpes de status de volta.',
    'mirror-armor': 'Reflete efeitos que reduzem atributos.',
    'synchronize': 'Devolve queimadura, paralisia e veneno a quem os causou.',
    'sturdy': 'Não é derrotado de uma vez com HP cheio.',
    'rough-skin': 'Machuca quem o atinge com contato.',
    'iron-barbs': 'Machuca quem o atinge com contato.',
    'aftermath': 'Machuca quem o derrota com contato.',
    'innards-out': 'Ao cair, causa dano igual ao HP que tinha.',
    'spicy-spray': 'Queima quem o atinge.',
    'toxic-debris': 'Espalha espinhos venenosos ao ser atingido.',
    'liquid-ooze': 'Quem rouba HP dele perde HP.',
    'perish-body': 'Quem o atinge com contato cai em três turnos.',
    'regenerator': 'Recupera 1/3 do HP ao sair de campo.',
    'hospitality': 'Ao entrar, recupera HP de um aliado.',
    # Ambiente
    'drought': 'Invoca sol forte ao entrar.',
    'desolate-land': 'Invoca sol extremo ao entrar.',
    'mega-sol': 'Usa os golpes como se houvesse sol forte.',
    'orichalcum-pulse': 'Invoca sol forte e aumenta o Ataque.',
    'drizzle': 'Invoca chuva ao entrar.',
    'primordial-sea': 'Invoca chuva forte ao entrar.',
    'sand-stream': 'Invoca tempestade de areia ao entrar.',
    'snow-warning': 'Invoca granizo ao entrar.',
    'electric-surge': 'Cria terreno elétrico ao entrar.',
    'hadron-engine': 'Cria terreno elétrico e aumenta o Ataque Especial.',
    'grassy-surge': 'Cria terreno de grama ao entrar.',
    'seed-sower': 'Cria terreno de grama ao ser atingido.',
    'psychic-surge': 'Cria terreno psíquico ao entrar.',
    'misty-surge': 'Cria terreno de névoa ao entrar.',
    'as-one-glastrier': 'Combina Unnerve (impede frutas) e Chilling Neigh (Ataque sobe ao derrotar).',
    'as-one-spectrier': 'Combina Unnerve (impede frutas) e Grim Neigh (Ataque Especial sobe ao derrotar).',
    'delta-stream': 'Cria uma corrente que anula as fraquezas dos Voadores.',
}


def main():
    wb = openpyxl.load_workbook(SRC)
    ws = wb['Efeitos (arquétipos)']
    head = [c.value for c in ws[1]]
    ci = {h: i for i, h in enumerate(head)}
    for row in ws.iter_rows(min_row=2):
        code = row[ci['Código']].value
        if code in RENAME:
            row[ci['Efeito']].value = RENAME[code]
    ws2 = wb['Habilidades']
    h2 = [c.value for c in ws2[1]]
    col = 'Na Pokédex (resumo)'
    if col not in h2:
        ws2.cell(row=1, column=len(h2) + 1, value=col)
        h2.append(col)
        ws2.column_dimensions[openpyxl.utils.get_column_letter(len(h2))].width = 60
    c2 = {h: i for i, h in enumerate(h2)}
    ename = {}
    for row in ws.iter_rows(min_row=2):
        ename[row[ci['Código']].value] = row[ci['Efeito']].value
    missing = set(NOTES)
    for row in ws2.iter_rows(min_row=2):
        slug = row[c2['Código']].value
        code = row[c2['Código do efeito']].value
        if code in RENAME:
            row[c2['Efeito']].value = RENAME[code]
        cell = ws2.cell(row=row[0].row, column=c2[col] + 1)
        if slug in NOTES:
            cell.value = NOTES[slug]
            missing.discard(slug)
    print('notas sem habilidade correspondente:', sorted(missing))
    wb.save(SRC)


main()
